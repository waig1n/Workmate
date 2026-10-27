from types import SimpleNamespace
from unittest.mock import AsyncMock

import pytest
from agentscope.message import TextBlock
from agentscope.rag import Chunk, VectorSearchResult, TextParser

from workmate_runtime.rag.components import NonEmptyParser
from workmate_runtime.rag.sources import source_projection


async def test_empty_documents_fail_instead_of_ready():
    with pytest.raises(ValueError, match="No extractable text"):
        await NonEmptyParser(TextParser()).parse(b"   ", "empty.txt")


def test_sources_are_stable_and_do_not_expose_internal_metadata():
    hit = VectorSearchResult(
        document_id="doc",
        score=0.8,
        chunk=Chunk(
            content=TextBlock(text="evidence"),
            source="report.pdf",
            chunk_index=2,
            total_chunks=3,
            metadata={"page": 3, "secret": "never"},
        ),
    )
    source = source_projection("product", hit)
    assert source == source_projection("product", hit)
    assert source["metadata"] == {"page": 3}
    assert source["document_id"] == "doc"
    assert source["knowledge_base_id"] == "product"


from fastapi import FastAPI
from httpx import AsyncClient, ASGITransport
from workmate_runtime.config import RuntimeSettings
from workmate_runtime.rag.management import KnowledgeManagement
from workmate_runtime.rag.management_router import create_management_router


def runtime_fixture(enabled=True):
    app = FastAPI()
    app.state.message_bus = SimpleNamespace(
        log_append=AsyncMock(return_value="source-entry"), publish=AsyncMock()
    )
    app.state.storage = SimpleNamespace(
        get_session=AsyncMock(
            return_value=SimpleNamespace(
                config=SimpleNamespace(knowledge_config=None), state=None
            )
        ),
        upsert_session=AsyncMock(),
        upsert_credential=AsyncMock(side_effect=lambda u, c: c.id),
    )
    app.state.knowledge_base_service = SimpleNamespace(
        _resolve_knowledge=AsyncMock(),
        _require_edit=AsyncMock(),
        get_document=AsyncMock(),
        register_document=AsyncMock(),
        delete_document=AsyncMock(),
        create_knowledge_base=AsyncMock(return_value={"id": "runtime-kb"}),
    )
    settings = RuntimeSettings(
        knowledge={"enabled": enabled},
        knowledge_embedding_model_type="openai",
        knowledge_embedding_model_name="embedding",
        knowledge_embedding_model_dimensions=1024,
        embedding_model_type="dashscope",
        embedding_model_name="reme-only",
        embedding_model_dimensions=768,
    )
    runtime = KnowledgeManagement(settings, app.state.storage, None, None, app.state.knowledge_base_service)
    async def resolve():
        return runtime
    app.include_router(create_management_router(resolve))
    return runtime, app


async def test_upload_guard_and_processing_delete_conflict():
    runtime, app = runtime_fixture()
    runtime.settings.knowledge.max_upload_bytes = 3
    async with AsyncClient(
        transport=ASGITransport(app=app),
        base_url="http://test",
        headers={"X-User-ID": "owner"},
    ) as client:
        assert (
            await client.put(
                "/workmate/knowledge/r/staged/d", files={"file": ("x.txt", b"1234")}
            )
        ).status_code == 413
        assert (
            await client.put(
                "/workmate/knowledge/r/staged/d", files={"file": ("x.html", b"x")}
            )
        ).status_code == 422
        app.state.knowledge_base_service.get_document.return_value = SimpleNamespace(
            status="indexing"
        )
        app.state.storage.get_knowledge_document = AsyncMock(
            return_value=SimpleNamespace(status="indexing")
        )
        assert (await client.delete("/workmate/knowledge/r/documents/d")).status_code == 409
        assert (
            await client.post("/workmate/knowledge/r/documents/d/retry")
        ).status_code == 409
    app.state.knowledge_base_service.register_document.assert_not_called()
    app.state.knowledge_base_service.delete_document.assert_not_called()


async def test_creation_pins_upstream_chunker_and_real_credential(monkeypatch):
    runtime, app = runtime_fixture()
    monkeypatch.setenv("KNOWLEDGE_EMBEDDING_MODEL_API_KEY", "test-secret")
    monkeypatch.setenv("EMBEDDING_MODEL_API_KEY", "reme-decoy")
    async with AsyncClient(
        transport=ASGITransport(app=app), base_url="http://test"
    ) as client:
        response = await client.post(
            "/workmate/knowledge", json={"name": "Manual"}, headers={"X-User-ID": "owner"}
        )
    assert response.status_code == 200, response.text
    credential = app.state.storage.upsert_credential.call_args.args[1]
    assert credential.api_key.get_secret_value() == "test-secret"
    args = app.state.knowledge_base_service.create_knowledge_base.call_args.args
    assert args[3].dimensions == 1024
    assert args[3].model == "embedding"
    assert args[3].type == "openai_credential"
    assert args[4].parameters == {"chunk_size": 512, "overlap": 64}


async def test_failed_retry_reuses_document_id_and_clears_vectors(monkeypatch):
    from agentscope.app.storage import KnowledgeDocumentRecord, KnowledgeDocumentData

    runtime, app = runtime_fixture()
    doc = KnowledgeDocumentRecord(
        id="d",
        user_id="owner",
        knowledge_base_id="r",
        status="error",
        data=KnowledgeDocumentData(
            filename="x.txt", size=3, blob_uri="s3://bucket/private", error="failure"
        ),
    )
    service = app.state.knowledge_base_service
    service.get_document.return_value = doc
    handle = SimpleNamespace(delete_document=AsyncMock())
    service._resolve_knowledge.return_value = handle
    app.state.storage.upsert_knowledge_document = AsyncMock(side_effect=lambda u, d: d)
    app.state.message_bus = object()
    enqueue = AsyncMock()
    monkeypatch.setattr("agentscope.app._bus_ops.enqueue_index_task", enqueue)
    async with AsyncClient(
        transport=ASGITransport(app=app), base_url="http://test"
    ) as client:
        response = await client.post(
            "/workmate/knowledge/r/documents/d/retry", headers={"X-User-ID": "owner"}
        )
    assert response.status_code == 200, response.text
    assert response.json()["id"] == "d"
    assert response.json()["status"] == "pending"
    assert "blob_uri" not in response.text
    handle.delete_document.assert_awaited_once_with("d")
    assert enqueue.call_args.kwargs["document_id"] == "d"


async def test_credential_rotation_preserves_pinned_identity(monkeypatch):
    from workmate_runtime.rag.components import refresh_embedding_credential

    runtime, app = runtime_fixture()
    app.state.storage.get_knowledge_base = AsyncMock(
        return_value=SimpleNamespace(
            data=SimpleNamespace(
                embedding_model_config=SimpleNamespace(credential_id="pinned")
            )
        )
    )
    app.state.storage.get_credential = AsyncMock(
        return_value=SimpleNamespace(
            data={
                "id": "pinned",
                "name": "embedding",
                "type": "openai_credential",
                "api_key": "old",
                "base_url": None,
            }
        )
    )
    monkeypatch.setenv("KNOWLEDGE_EMBEDDING_MODEL_API_KEY", "rotated")
    monkeypatch.setenv("EMBEDDING_MODEL_API_KEY", "reme-decoy")
    await refresh_embedding_credential(
        runtime.settings, app.state.storage, "owner", "kb"
    )
    credential = app.state.storage.upsert_credential.call_args.args[1]
    assert credential.id == "pinned"
    assert credential.api_key.get_secret_value() == "rotated"
    app.state.storage.upsert_credential.reset_mock()
    runtime.settings.knowledge_embedding_model_base_url = "https://new-endpoint.invalid"
    await refresh_embedding_credential(
        runtime.settings, app.state.storage, "owner", "kb"
    )
    app.state.storage.upsert_credential.assert_not_awaited()


def test_wrapped_parser_preserves_extension_catalog():
    assert ".txt" in NonEmptyParser(TextParser()).supported_extensions()


async def test_common_office_parsers_preserve_text_and_sheet_location():
    import io
    from docx import Document
    from pptx import Presentation
    from openpyxl import Workbook
    from agentscope.rag import WordParser, PPTParser, ExcelParser

    word = Document()
    word.add_paragraph("Employee handbook evidence")
    word_bytes = io.BytesIO()
    word.save(word_bytes)
    ppt = Presentation()
    slide = ppt.slides.add_slide(ppt.slide_layouts[1])
    slide.shapes.title.text = "Quarterly evidence"
    ppt_bytes = io.BytesIO()
    ppt.save(ppt_bytes)
    book = Workbook()
    book.active.title = "Finance"
    book.active.append(["Item", "Amount"])
    book.active.append(["Revenue", 100])
    xlsx_bytes = io.BytesIO()
    book.save(xlsx_bytes)
    for parser, content, filename in [
        (WordParser(), word_bytes.getvalue(), "guide.docx"),
        (PPTParser(), ppt_bytes.getvalue(), "report.pptx"),
        (ExcelParser(separate_sheet=True), xlsx_bytes.getvalue(), "report.xlsx"),
    ]:
        sections = await NonEmptyParser(parser).parse(content, filename)
        assert sections
        assert all(isinstance(s.content, TextBlock) for s in sections)
        if filename.endswith(".xlsx"):
            assert sections[0].metadata["sheet"] == "Finance"


async def test_real_qdrant_rustfs_component_construction():
    from workmate_runtime.rag.components import build_knowledge_components

    runtime, app = runtime_fixture()
    runtime.settings.knowledge_qdrant_url = "http://127.0.0.1:6333"
    runtime.settings.knowledge_qdrant_api_key = "test-key"
    runtime.settings.knowledge_s3_endpoint = "http://127.0.0.1:9000"
    runtime.settings.knowledge_s3_bucket = "knowledge"
    runtime.settings.knowledge_s3_access_key = "test"
    runtime.settings.knowledge_s3_secret_key = "test"
    parts = build_knowledge_components(runtime.settings, app.state.storage)
    assert len(parts["knowledge_parsers"]) == 5
    assert parts["enable_index_worker"] is True
    assert all(p.supported_extensions() for p in parts["knowledge_parsers"])
    await parts["knowledge_base_manager"].__aexit__(None, None, None)


def test_qdrant_component_refuses_unauthenticated_backend():
    from workmate_runtime.rag.components import build_knowledge_components

    runtime, app = runtime_fixture()
    with pytest.raises(ValueError, match="KNOWLEDGE_QDRANT_API_KEY"):
        build_knowledge_components(runtime.settings, app.state.storage)


@pytest.mark.parametrize("change", [None, "type", "model", "dimensions", "endpoint"])
async def test_legacy_dimension_override_requires_matching_pinned_identity(change):
    from workmate_runtime.rag.components import apply_embedding_request_policy

    runtime, app = runtime_fixture()
    runtime.settings.knowledge_embedding_model_pass_dimensions = False
    config = SimpleNamespace(
        type="openai_credential",
        model="embedding",
        dimensions=1024,
        credential_id="c",
        parameters={},
    )
    data = {"type": "openai_credential", "base_url": None}
    if change == "endpoint":
        data["base_url"] = "https://other.example"
    elif change:
        setattr(
            config,
            change,
            {"type": "dashscope_credential", "model": "other", "dimensions": 768}[
                change
            ],
        )
    app.state.storage.get_knowledge_base = AsyncMock(
        return_value=SimpleNamespace(
            data=SimpleNamespace(embedding_model_config=config)
        )
    )
    app.state.storage.get_credential = AsyncMock(
        return_value=SimpleNamespace(data=data)
    )
    model = SimpleNamespace(pass_dimensions=True, dimensions=1024)
    await apply_embedding_request_policy(
        runtime.settings,
        app.state.storage,
        "owner",
        "kb",
        SimpleNamespace(embedding_model=model),
    )
    assert model.pass_dimensions is (change is not None)
    assert model.dimensions == 1024
    assert config.parameters == {}


@pytest.mark.parametrize("pinned", [True, False])
async def test_persisted_dimension_policy_controls_actual_upstream_request(pinned):
    from workmate_runtime.rag.components import apply_embedding_request_policy
    from agentscope.embedding import OpenAIEmbeddingModel
    from agentscope.credential import OpenAICredential

    runtime, app = runtime_fixture()
    runtime.settings.knowledge_embedding_model_pass_dimensions = not pinned
    config = SimpleNamespace(
        type="openai_credential",
        model="old-model",
        dimensions=1024,
        credential_id="c",
        parameters={"pass_dimensions": pinned},
    )
    app.state.storage.get_knowledge_base = AsyncMock(
        return_value=SimpleNamespace(
            data=SimpleNamespace(embedding_model_config=config)
        )
    )
    model = OpenAIEmbeddingModel(
        credential=OpenAICredential(api_key="unused"),
        model="old-model",
        dimensions=1024,
        parameters=OpenAIEmbeddingModel.Parameters(**config.parameters),
    )
    real_client = model.client
    create = AsyncMock(
        return_value=SimpleNamespace(
            data=[SimpleNamespace(embedding=[0.0] * 1024)],
            usage=SimpleNamespace(total_tokens=1),
        )
    )
    model.client = SimpleNamespace(embeddings=SimpleNamespace(create=create))
    try:
        await apply_embedding_request_policy(
            runtime.settings,
            app.state.storage,
            "owner",
            "kb",
            SimpleNamespace(embedding_model=model),
        )
        result = await model(["synthetic text"])
        assert ("dimensions" in create.call_args.kwargs) is pinned
        assert "pass_dimensions" not in create.call_args.kwargs
        assert len(result.embeddings[0]) == 1024
        assert model.dimensions == 1024
    finally:
        await real_client.close()


def test_index_error_projection_distinguishes_provider_rejection_without_leaking_details():
    from agentscope.app.storage import KnowledgeDocumentRecord, KnowledgeDocumentData
    from workmate_runtime.rag.document_projection import document_projection

    doc = KnowledgeDocumentRecord(
        id="doc",
        user_id="owner",
        knowledge_base_id="kb",
        status="error",
        data=KnowledgeDocumentData(
            filename="sample.pdf",
            size=1,
            blob_uri="s3://private",
            error="BadRequestError: provider secret-url private-key",
        ),
    )
    assert (
        document_projection(doc)["error"]
        == "Embedding request rejected; check model parameters and provider compatibility"
    )
    doc.data.error = "UnexpectedError: private-key"
    message = document_projection(doc)["error"]
    assert "diagnostic" in message.lower()
    assert "private-key" not in message
    assert "format" not in message


@pytest.mark.parametrize(
    "provider,expected", [("openai", {"pass_dimensions": False}), ("dashscope", {})]
)
async def test_creation_persists_openai_wire_option_only(
    provider, expected, monkeypatch
):
    runtime, app = runtime_fixture()
    runtime.settings.knowledge_embedding_model_type = provider
    runtime.settings.knowledge_embedding_model_pass_dimensions = False
    monkeypatch.setenv("KNOWLEDGE_EMBEDDING_MODEL_API_KEY", "test-only")
    async with AsyncClient(
        transport=ASGITransport(app=app), base_url="http://test"
    ) as client:
        response = await client.post(
            "/workmate/knowledge", json={"name": "Example"}, headers={"X-User-ID": "owner"}
        )
    assert response.status_code == 200, response.text
    config = app.state.knowledge_base_service.create_knowledge_base.call_args.args[3]
    assert config.parameters == expected
    assert config.dimensions == 1024


async def test_manager_applies_persisted_wire_policy_after_upstream_construction(
    monkeypatch,
):
    from agentscope.app.rag.knowledge_base_manager import CollectionPerKbManager
    from workmate_runtime.rag.components import build_knowledge_components

    runtime, app = runtime_fixture()
    runtime.settings.knowledge_qdrant_url = "http://127.0.0.1:6333"
    runtime.settings.knowledge_qdrant_api_key = "test-key"
    runtime.settings.knowledge_s3_endpoint = "http://127.0.0.1:9000"
    runtime.settings.knowledge_s3_bucket = "test"
    runtime.settings.knowledge_s3_access_key = "test"
    runtime.settings.knowledge_s3_secret_key = "test"
    config = SimpleNamespace(
        type="openai_credential",
        model="embedding",
        dimensions=1024,
        credential_id="c",
        parameters={"pass_dimensions": False},
    )
    app.state.storage.get_knowledge_base = AsyncMock(
        return_value=SimpleNamespace(
            data=SimpleNamespace(embedding_model_config=config)
        )
    )
    app.state.storage.get_credential = AsyncMock(return_value=None)
    handle = SimpleNamespace(
        embedding_model=SimpleNamespace(pass_dimensions=True, dimensions=1024)
    )
    monkeypatch.setattr(
        CollectionPerKbManager, "get_knowledge", AsyncMock(return_value=handle)
    )
    manager = build_knowledge_components(runtime.settings, app.state.storage)[
        "knowledge_base_manager"
    ]
    try:
        assert await manager.get_knowledge("owner", "kb") is handle
        assert handle.embedding_model.pass_dimensions is False
        assert handle.embedding_model.dimensions == 1024
    finally:
        await manager.__aexit__(None, None, None)
